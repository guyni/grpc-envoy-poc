#!/usr/bin/env python3
"""Build docs/presentation.pptx from the POC results.

Run with a python that has python-pptx:
    python3 -m venv /tmp/deckvenv && /tmp/deckvenv/bin/pip install python-pptx
    /tmp/deckvenv/bin/python scripts/build-deck.py

The deck is 16:9, uses native shapes and charts so Google Slides imports them
as editable objects (File -> Import slides), and every slide carries speaker notes.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

NAVY = RGBColor(0x1F, 0x2A, 0x44)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
GRAY = RGBColor(0x4B, 0x55, 0x63)
LIGHT = RGBColor(0xF2, 0xF3, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xB4, 0x23, 0x18)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def text(s, left, top, width, height, value, size=18, bold=False, color=GRAY,
         align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = value
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Calibri"
    p.alignment = align
    return tb


def title(s, value, sub=None):
    text(s, 0.8, 0.4, 11.7, 0.9, value, size=30, bold=True, color=NAVY)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.18),
                             Inches(1.3), Inches(0.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    if sub:
        text(s, 0.8, 1.32, 11.7, 0.5, sub, size=15)


def bullets(s, items, left=0.9, top=1.85, width=11.4, height=4.9, size=19):
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        level = 0
        if isinstance(item, tuple):
            item, level = item
        r = p.add_run()
        r.text = ("•  " if level == 0 else "–  ") + item
        r.font.size = Pt(size if level == 0 else size - 3)
        r.font.name = "Calibri"
        r.font.color.rgb = NAVY if level == 0 else GRAY
        p.level = level


def notes(s, value):
    s.notes_slide.notes_text_frame.text = value


def table(s, data, left, top, width, height, font=14, widths=None):
    shape = s.shapes.add_table(len(data), len(data[0]), Inches(left), Inches(top),
                               Inches(width), Inches(height))
    tbl = shape.table
    if widths:
        for i, w in enumerate(widths):
            tbl.columns[i].width = Inches(w)
    for r, row in enumerate(data):
        for c, cell_value in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = str(cell_value)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(font)
                    run.font.name = "Calibri"
                    run.font.bold = (r == 0)
                    run.font.color.rgb = WHITE if r == 0 else NAVY
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (LIGHT if r % 2 == 0 else WHITE)
    return tbl


def box(s, left, top, width, height, value, fill=WHITE, color=NAVY, size=14, bold=True):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
                             Inches(width), Inches(height))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = color
    tf = shp.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = value
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Calibri"
    p.alignment = PP_ALIGN.CENTER


def arrow(s, left, top, width, height, label=None):
    shp = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(left), Inches(top),
                             Inches(width), Inches(height))
    shp.fill.solid()
    shp.fill.fore_color.rgb = TEAL
    shp.line.fill.background()
    if label:
        text(s, left - 0.5, top - 0.62, width + 1.0, 0.5, label, size=11,
             color=TEAL, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- slide 1
s = slide()
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()
text(s, 0.9, 2.2, 11.5, 1.4, "Reducing msgbroker TCP connections\nwith an Envoy relay", size=40, bold=True, color=WHITE)
text(s, 0.9, 3.9, 11.5, 0.6, "POC findings and a deployment recommendation", size=22, color=RGBColor(0xB8, 0xC4, 0xD6))
text(s, 0.9, 5.6, 11.5, 0.5, "github.com/guyni/grpc-envoy-poc", size=14, color=RGBColor(0x8E, 0x9B, 0xB0))
notes(s, "Open with the bottom line: we built the relay from the proposal, measured it from 5 to 1000 agents, and the mechanism works - but the numbers say it should not be deployed at our current fleet size. Everything after this is the evidence.")

# ---------------------------------------------------------------- slide 2
s = slide()
title(s, "The problem: one socket per agent")
bullets(s, [
    "Each node agent holds one long-lived bidirectional gRPC stream to msgbroker",
    "That stream is never idle-closed: it lives as long as the agent",
    "So the endpoint's TCP socket count tracks fleet size, not traffic",
    "The concern is growth, not today's load",
    ("1000 agents per cluster across 100 clusters = 100k sockets before traffic shapes them", 1),
])
notes(s, "The proposal's first instruction is not 'build the relay' but 'measure which resource runs out first'. That is why we ran the baseline before recommending anything.")

# ---------------------------------------------------------------- slide 3
s = slide()
title(s, "The proposal: one relay per target cluster")
box(s, 0.6, 2.4, 2.1, 1.3, "Node agents\nx N")
arrow(s, 2.8, 2.85, 1.1, 0.4, "mTLS + token\nN connections")
box(s, 4.0, 2.25, 2.4, 1.6, "Envoy relay\n:8443\n1 worker")
arrow(s, 6.5, 2.85, 1.1, 0.4, "pooled HTTP/2\nTLS")
box(s, 7.7, 2.4, 2.1, 1.3, "Traefik\n:443")
arrow(s, 9.9, 2.85, 1.1, 0.4, "h2c\n1 socket")
box(s, 11.1, 2.4, 2.1, 1.3, "msgbroker\npod")
bullets(s, [
    "The RPC contract, the per-agent token and one authenticated stream per agent are unchanged",
    "The relay runs on the host, independent of the Kubernetes API and CNI",
    "Cost: a shared failure domain - one relay per cluster",
], top=4.5, height=2.2, size=17)
notes(s, "Two things make this attractive: no protobuf or service changes, and identity survives because the token travels in each stream's metadata even when TCP connections are shared.")

# ---------------------------------------------------------------- slide 4
s = slide()
title(s, "What we set out to verify")
table(s, [
    ["Claim", "How it is proven"],
    ["C1  Pooling", "N agents produce far fewer connections at the public endpoint"],
    ["C2  Isolation", "RPC metadata survives pooling; no cross-agent delivery"],
    ["C3  Admission", "Agent mTLS plus per-RPC token, failing closed"],
    ["C4  Recovery", "Bounded detection and reconnect after realistic faults"],
    ["C5  Scale", "Behaviour holds at 5, 10, 200 and 1000 agents"],
], left=0.9, top=1.8, width=11.5, height=3.9, widths=[2.8, 8.7])
text(s, 0.9, 6.1, 11.5, 0.8, "Out of scope: production token issuance, multi-cluster capacity, a real control-plane host with systemd.",
     size=14)
notes(s, "We validate the mechanism and its operational shape, not a production rollout.")

# ---------------------------------------------------------------- slide 5
s = slide()
title(s, 'What we built')
bullets(s, [
    'Dedicated kind cluster: cert-manager CA issuer + Traefik, created by one script',
    'msgbroker stand-in: real bidi gRPC service, per-agent tokens, session generations',
    'Go agents: one ClientConn, certificate and token per simulated agent',
    'Envoy relay: container on the host side, mTLS in, pooled HTTP/2 out, one worker',
    'Harness: measured gauges, 12-case auth matrix, fault injection',
])
notes(s, 'Synthetic stand-in for the msgbroker contract, not a reimplementation: one long-lived stream per agent, server-initiated commands, acks, heartbeats. The whole environment is reproducible from one command sequence.')

# ---------------------------------------------------------------- slide 6
s = slide()
title(s, 'The measurement that matters')
bullets(s, [
    'Metric: TCP connections into the gRPC endpoint grpc.localtest.me:443',
    'Read from Traefik metrics: traefik_open_connections on the websecure entrypoint',
    'Sampled every second during each run - measured, not inferred from agent count',
    'Cross-checked against Envoy downstream/upstream gauges and msgbroker accepted sockets',
    'Two paths, identical agents: P1 direct, P2 through the relay',
], size=18)
notes(s, 'If someone asks how we counted: three independent observers, same numbers. That is the defence of this whole deck.')

# ---------------------------------------------------------------- slide 7
s = slide()
title(s, 'Headline result', 'Connections into grpc.localtest.me:443')
table(s, [
    ['Agents', 'Relay upstream pool', 'P1 - no relay', 'P2 - relay'],
    ['5', '1', '5', '1'],
    ['10', '1', '10', '1'],
    ['200', '1', '200', '1'],
    ['1000', '4', '1000', '4'],
], left=0.9, top=1.9, width=9.2, height=3.4, widths=[1.8, 2.6, 2.4, 2.4])
box(s, 10.4, 2.2, 2.4, 1.1, '1000 → 4', fill=LIGHT, color=TEAL, size=26)
text(s, 10.2, 3.4, 2.8, 0.8, 'at the public endpoint', size=12, align=PP_ALIGN.CENTER)
bullets(s, ['Every command acked, zero cross-agent deliveries, zero reconnects'], top=5.6, height=1.0, size=17)
notes(s, 'Read the table out loud, then the sentence the proposal cares about: without the relay the endpoint holds one connection per agent; with it, the relay pool. Growth from 1 to 4 at 1000 agents is the peer stream cap being reached, not a bug.')

# ---------------------------------------------------------------- slide 8
s = slide()
title(s, 'How the pooling behaves', 'Same 1000 streams, two endpoint connection counts')
text(s, 1.2, 1.9, 5.0, 0.4, 'P1 - direct', size=15, bold=True, color=GRAY)
box(s, 1.2, 2.5, 1.7, 0.9, 'agents', size=13)
arrow(s, 3.0, 2.75, 1.5, 0.4, '1000 conns')
box(s, 4.6, 2.5, 1.7, 0.9, 'Traefik', size=13)
text(s, 7.4, 1.9, 5.0, 0.4, 'P2 - through the relay', size=15, bold=True, color=GRAY)
box(s, 7.4, 2.5, 1.4, 0.9, 'agents', size=13)
arrow(s, 8.9, 2.75, 0.9, 0.4, '1000')
box(s, 9.9, 2.45, 1.5, 1.0, 'relay', size=13)
arrow(s, 11.5, 2.75, 0.8, 0.4, '1 or 4')
box(s, 12.3, 2.5, 0.9, 0.9, 'Traefik', size=9)
bullets(s, [
    'One connection pool per Envoy worker: one worker gives one upstream connection',
    'HTTP/2 multiplexes streams onto it up to the peer cap (~250 here)',
    'Envoy opens more only when a connection is full - there is no pool-size setting',
    '1 connection up to 200 agents; 4 at 1000 = ceil(1000 / 250)',
], top=4.3, height=2.4, size=17)
notes(s, 'Envoy docs describe the mechanism, not a connection-reduction feature. The reduction is emergent: multiplexing plus the peer stream cap plus the worker count. Config surface is max_concurrent_streams, circuit-breaker ceilings and concurrency.')

# ---------------------------------------------------------------- slide 9
s = slide()
title(s, 'The twist: the ingress already pools')
box(s, 0.9, 1.9, 3.1, 1.5, '1 socket\nat msgbroker', fill=LIGHT, color=RED, size=20)
text(s, 4.2, 1.9, 8.2, 1.5, 'on both paths, at every agent count - that is Traefik pooling, present before the relay existed', size=17)
bullets(s, [
    'So the relay reduces connections at the public endpoint, not at msgbroker',
    'At fewer than ~250 streams it adds a hop and a failure domain for no backend benefit',
    "This is the proposal's own caveat in section 1, answered by measurement",
    'Anyone sizing the relay must measure where the real bottleneck is first',
], top=3.7, height=3.0, size=18)
notes(s, 'This is the finding that changes the recommendation. Our baseline answered the proposal question: yes, the public endpoint already multiplexes backend connections.')

# ---------------------------------------------------------------- slide 10
s = slide()
title(s, 'Latency: better small, worse in the tail')
data = CategoryChartData()
data.categories = ['5', '10', '200', '1000']
data.add_series('P1 - direct', (2.08, 2.72, 1.33, 1.18))
data.add_series('P2 - relay', (0.42, 0.93, 8.07, 9.34))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.9), Inches(1.8),
                        Inches(7.8), Inches(4.7), data)
chart = gf.chart
chart.has_title = True
chart.chart_title.text_frame.text = 'Command latency p99 (ms) by agent count'
chart.has_legend = True
chart.legend.position = XL_LEGEND_POSITION.BOTTOM
chart.legend.include_in_layout = False
plot = chart.plots[0]
plot.has_data_labels = True
plot.data_labels.font.size = Pt(10)
bullets(s, [
    'Small N: fewer TLS handshakes, one pooled hop - relay is faster',
    "At scale: 1000 streams share one worker's pools and flow-control windows",
    'Same-host loopback, so this isolates relay cost without WAN effects',
], left=9.0, top=2.0, width=3.9, height=4.4, size=14)
notes(s, 'The proposal warns that a shared pooled connection shares congestion and flow control. The 1000-agent p99 is the local evidence of that trade-off.')

# ---------------------------------------------------------------- slide 11
s = slide()
title(s, 'Authentication and isolation: 12 of 12 fail closed')
bullets(s, [
    'Agent mTLS (rejected at handshake)',
    ('no certificate, wrong-cluster SAN, expired, untrusted CA', 1),
    'Per-RPC token (Unauthenticated)',
    ('missing, expired, forged signature, wrong audience, wrong issuer, no identity', 1),
    'Identity binding',
    ('Hello identity must match the token cluster and node, or PermissionDenied', 1),
    'Envoy counters separate the causes: fail_verify_no_cert, fail_verify_san, fail_verify_error',
], size=18)
notes(s, 'The point: sharing a TCP connection does not merge identity. Each stream still carries its own token, and the server binds the session to the token claims, not the TLS peer or source IP. No negative case reached a registered session.')

# ---------------------------------------------------------------- slide 12
s = slide()
title(s, 'Failure behaviour at 5 agents')
table(s, [
    ['Event', 'Detection', 'Recovery'],
    ['Upstream TCP reset', 'immediate', '2 s'],
    ['Relay restart', 'immediate', '2 s'],
    ['App hang, transport healthy', '10 s (heartbeat)', '1 s after release'],
    ['Silent blackhole', '40 s (TCP_USER_TIMEOUT)', '25 s after restore'],
], left=0.9, top=1.9, width=11.4, height=3.2, widths=[5.0, 3.4, 3.0])
bullets(s, [
    'A live stream is not a live session: if the app stalls while Hello is in flight, the stream stays up but the agent is unregistered',
    'Fix: require Welcome within the heartbeat timeout, or reconnect',
], top=5.3, height=1.7, size=16)
notes(s, 'Two detectors: the 40 s socket option caught the blackhole before the 60 s HTTP/2 PING could; the heartbeat is the only thing that catches a hung application on a healthy transport. All recoveries are inside the 120 s bar.')

# ---------------------------------------------------------------- slide 13
s = slide()
title(s, 'What surprised us')
bullets(s, [
    'Traefik IngressRoute needs the named service port; a numeric port returns 503',
    'The core-Ingress h2c annotation does not take effect here; the CRD with scheme h2c does',
    'Envoy v1.32 rejects two settings from the proposal: stream_flush_timeout, and resource_api_version V3 in YAML',
    'envoy --mode validate cannot resolve filesystem SDS - a green validate proves less than it looks',
    'LOGICAL_DNS needs the getaddrinfo resolver, or the container ignores /etc/hosts',
], size=17)
notes(s, 'Small but expensive to rediscover. All of them are written down in the results doc with the exact symptom.')

# ---------------------------------------------------------------- slide 14
s = slide()
title(s, 'What we did not prove')
bullets(s, [
    'A real control-plane host: systemd, host DNS, host networking during a Kubernetes outage',
    'One-direction packet loss - only a full blackhole was measured',
    'A 24-hour soak; the longest run is 1000 agents for 60 s',
    'Multi-relay or multi-cluster capacity, production token issuance, CRL distribution',
    'Latency beyond same-host loopback',
])
notes(s, 'The first two bullets carry engineering work. Everything else is a deliberate scope line, not a finding.')

# ---------------------------------------------------------------- slide 15
s = slide()
title(s, 'Recommendation')
box(s, 0.9, 1.85, 11.5, 0.95, "Do not deploy the relay at today's fleet size (under 10 agents per cluster)",
    fill=LIGHT, color=RED, size=18)
bullets(s, [
    'It reduces endpoint connections, but adds a hop, a CA, a host process and a shared failure domain',
    'msgbroker socket count is unchanged, because the ingress already pools',
    'The relay is justified only when the public endpoint socket count is the measured bottleneck',
    'At that point: measure the real endpoint stream cap and backend pooling first - the answer may be a limit increase, not a relay',
], top=3.2, height=3.4, size=17)
notes(s, 'Slow down here. The relay works; the question is whether it buys anything at our scale. At 200+ agents the endpoint reduction is real (200 to 1), and at 1000 the pool grows to 4 - still a 250x reduction at the endpoint.')

# ---------------------------------------------------------------- slide 16
s = slide()
title(s, 'Decision requested')
bullets(s, [
    'Decide on a pilot, not a rollout - and only after the proposal section 1 baseline measurements',
    'Baseline to capture: accepted sockets at the endpoint and each backend hop, peer stream caps, fds and CPU per connection',
    'Add the missing failure case (one-direction loss) and a host deployment model before any pilot',
    'If the endpoint turns out to be socket-bound, this POC says the relay will do the job',
])
notes(s, 'Frame it as risk reduction: a few days of work established that the ingress already pools, which saves a fleet-wide rollout of a component that would not have reduced msgbroker load at our scale.')

# ---------------------------------------------------------------- slide 17
s = slide()
title(s, 'Reproduce / live demo')
bullets(s, [
    'Environment in one sequence: 00-cluster-up, 20-deploy, 30-certs, 40-relay-up',
    'Compare paths: 50-run-path.sh --path p1|p2 --agents N --duration 20s',
    'Live gauges: relay admin 127.0.0.1:9901/stats, msgbroker /stats, Traefik open connections',
    'Evidence: out/runs, out/auth-matrix/results.md, out/faults',
], size=18)
box(s, 0.9, 5.3, 11.5, 0.9, 'Live demo: 10 agents on p1 then p2 - the ingress gauge goes 10 to 1 in about 40 seconds', fill=LIGHT, color=TEAL, size=16)
notes(s, 'If you demo live, two terminals, ten agents each path. That single number is the whole talk.')

# ---------------------------------------------------------------- slide 18
s = slide()
title(s, 'Summary')
bullets(s, [
    'The mechanism works: endpoint connections 5/10/200/1000 become 1/1/1/4, streams unchanged',
    'The ingress already pools backend connections, so at under 10 agents the relay adds cost, not value',
    'Auth and recovery behave as required: 12 of 12 closed, all recoveries inside 120 s',
    'Recommendation: no deployment now; revisit if the endpoint is measured to be socket-bound',
], size=19)
text(s, 0.9, 5.9, 11.5, 0.6, 'The relay is sound engineering for a problem we may not have.', size=20, bold=True, color=TEAL)
notes(s, 'Close with the sentence you opened with, now earned. Likely questions: what would make you deploy it (a measured endpoint socket bottleneck), and what did you not test (one-direction loss, real host outage, soak).')

prs.save('docs/presentation.pptx')
print('wrote docs/presentation.pptx')
